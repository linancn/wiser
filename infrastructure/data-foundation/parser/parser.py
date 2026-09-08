"""Bounded extraction of source content. Source files never supply executable code."""

import hashlib
import math
import re
import stat
import zipfile
from datetime import date, datetime, time, timedelta
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath

MAX_INPUT_BYTES = 64 * 1024 * 1024
MAX_EXPANDED_BYTES = 256 * 1024 * 1024
MAX_RECORDS = 2_000_000
MAX_TEXT = 1_000_000


class ParseError(Exception):
    pass


def safe_archive(path):
    archive = zipfile.ZipFile(path)
    try:
        entries = archive.infolist()
        if (
            len(entries) > 4096
            or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES
        ):
            raise ParseError("ARCHIVE_LIMIT")
        seen = set()
        for entry in entries:
            name = entry.orig_filename
            parts = PurePosixPath(name).parts
            if (
                not name
                or len(name) > 1024
                or "\\" in name
                or "\0" in name
                or name.startswith("/")
                or ".." in parts
                or re.match(r"^[a-zA-Z]:", name)
                or name in seen
                or stat.S_ISLNK(entry.external_attr >> 16)
            ):
                raise ParseError("UNSAFE_ARCHIVE")
            if entry.flag_bits & 1:
                raise ParseError("ENCRYPTED_CONTENT")
            if entry.file_size > max(1, entry.compress_size) * 1000:
                raise ParseError("ARCHIVE_LIMIT")
            seen.add(name)
        return archive
    except BaseException:
        archive.close()
        raise


def schema(labels, extra=()):
    if len(labels) + len(extra) > 256:
        raise ParseError("COLUMN_LIMIT")
    return {
        "type": "schema",
        "columns": [
            {"key": f"c{index + 1}", "label": str(label)[:512] or f"Column {index + 1}"}
            for index, label in enumerate(labels)
        ]
        + list(extra),
    }


def record(values, geometry=None, source_id=None, source_crs=None):
    return {
        "type": "record",
        "values": values,
        "geometry": geometry,
        "sourceId": source_id,
        "sourceCrs": source_crs,
    }


def cell_value(cell, epoch):
    value = cell.value
    if cell.data_type == "f":
        return {"formula": str(value)}
    if cell.data_type == "e":
        return {"error": str(value)}
    if isinstance(value, (datetime, date, time)):
        return {
            "localDateTime": value.isoformat(),
            "numberFormat": cell.number_format,
            "excelEpoch": epoch.isoformat(),
        }
    if isinstance(value, timedelta):
        return {
            "durationSeconds": value.total_seconds(),
            "numberFormat": cell.number_format,
        }
    if isinstance(value, float) and not math.isfinite(value):
        raise ParseError("INVALID_CONTENT")
    return value


def workbook(path):
    from openpyxl import load_workbook

    if not zipfile.is_zipfile(path):
        raise ParseError("INVALID_FORMAT")
    with safe_archive(path):
        pass
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheets = []
        labels = []
        for sheet in book.worksheets:
            sheet.reset_dimensions()
            first = next(
                (
                    row
                    for row in sheet.iter_rows()
                    if any(cell.value is not None for cell in row)
                ),
                (),
            )
            if len(first) > 254:
                raise ParseError("COLUMN_LIMIT")
            values = [cell.value for cell in first]
            has_header = (
                bool(values)
                and all(isinstance(value, str) and value.strip() for value in values)
                and len(set(values)) == len(values)
            )
            headers = (
                values
                if has_header
                else [f"Column {index + 1}" for index in range(len(values))]
            )
            for label in headers:
                if label not in labels:
                    labels.append(label)
            row_number = next((cell.row for cell in first if cell.value is not None), 1)
            sheets.append((sheet, headers, row_number + int(has_header)))
        yield schema(
            labels,
            (
                {"key": "__sheet", "label": "Sheet"},
                {"key": "__row", "label": "Source row"},
            ),
        )
        for sheet, headers, start in sheets:
            for cells in sheet.iter_rows(min_row=start):
                if not any(cell.value is not None for cell in cells):
                    continue
                if len(cells) > len(headers):
                    raise ParseError("INCONSISTENT_COLUMNS")
                values = {f"c{index + 1}": None for index in range(len(labels))}
                for index, cell in enumerate(cells):
                    values[f"c{labels.index(headers[index]) + 1}"] = cell_value(
                        cell, book.epoch
                    )
                values["__sheet"] = sheet.title
                values["__row"] = next(
                    cell.row for cell in cells if cell.value is not None
                )
                yield record(values)
    finally:
        book.close()


def legacy_workbook(path):
    import xlrd

    if path.read_bytes()[:8] != bytes.fromhex("d0cf11e0a1b11ae1"):
        raise ParseError("INVALID_FORMAT")
    book = xlrd.open_workbook(path, on_demand=True)
    try:
        sheets = []
        labels = []
        for sheet in book.sheets():
            if sheet.ncols > 254:
                raise ParseError("COLUMN_LIMIT")
            first = sheet.row_values(0) if sheet.nrows else []
            has_header = (
                bool(first)
                and all(isinstance(value, str) and value.strip() for value in first)
                and len(set(first)) == len(first)
            )
            headers = (
                first
                if has_header
                else [f"Column {index + 1}" for index in range(sheet.ncols)]
            )
            for label in headers:
                if label not in labels:
                    labels.append(label)
            sheets.append((sheet, headers, int(has_header)))
        yield schema(
            labels,
            (
                {"key": "__sheet", "label": "Sheet"},
                {"key": "__row", "label": "Source row"},
            ),
        )
        for sheet, headers, start in sheets:
            for row_index in range(start, sheet.nrows):
                cells = sheet.row(row_index)
                if all(
                    cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
                    for cell in cells
                ):
                    continue
                values = {f"c{index + 1}": None for index in range(len(labels))}
                for index, cell in enumerate(cells):
                    value = cell.value
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        value = {"excelSerial": value, "dateSystem": book.datemode}
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = {"error": str(value)}
                    values[f"c{labels.index(headers[index]) + 1}"] = value
                values["__sheet"] = sheet.name
                values["__row"] = row_index + 1
                yield record(values)
    finally:
        book.release_resources()


class VisibleText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = []

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "template"):
            self.hidden.append(tag)
        if not self.hidden and tag in (
            "p",
            "div",
            "br",
            "h1",
            "h2",
            "h3",
            "tr",
            "li",
            "section",
        ):
            self.parts.append("\n\n")

    def handle_endtag(self, tag):
        if self.hidden and tag == self.hidden[-1]:
            self.hidden.pop()
        if not self.hidden and tag in (
            "p",
            "div",
            "h1",
            "h2",
            "h3",
            "tr",
            "li",
            "section",
        ):
            self.parts.append("\n\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def document(path, kind):
    yield schema(["Text", "Source location"])
    if kind == "pdf":
        from pypdf import PdfReader

        if path.read_bytes()[:5] != b"%PDF-":
            raise ParseError("INVALID_FORMAT")
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ParseError("ENCRYPTED_CONTENT")
        if len(reader.pages) > 2000:
            raise ParseError("PAGE_LIMIT")
        segments = (
            (page.extract_text() or "", f"page:{index + 1}")
            for index, page in enumerate(reader.pages)
        )
    else:
        text = path.read_text(encoding="utf-8-sig")
        if kind == "html":
            parser = VisibleText()
            parser.feed(text)
            parser.close()
            text = "".join(parser.parts)
        segments = ((text, "document"),)
    total = 0
    for text, location in segments:
        for paragraph in re.split(r"\n\s*\n", text):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            total += len(paragraph)
            if total > MAX_TEXT:
                yield {"type": "warning", "reason": "TEXT_LIMIT"}
                return
            for start in range(0, len(paragraph), 16000):
                yield record({"c1": paragraph[start : start + 16000], "c2": location})
    if kind == "pdf" and total == 0:
        yield {"type": "warning", "reason": "TEXT_UNAVAILABLE"}


def archive_inventory(path):
    yield schema(["Member path", "Bytes", "SHA-256"])
    with safe_archive(path) as archive:
        for entry in archive.infolist():
            if entry.is_dir():
                continue
            digest = hashlib.sha256()
            size = 0
            with archive.open(entry) as member:
                for chunk in iter(lambda: member.read(65536), b""):
                    size += len(chunk)
                    if size > entry.file_size or size > MAX_EXPANDED_BYTES:
                        raise ParseError("ARCHIVE_LIMIT")
                    digest.update(chunk)
            yield record({"c1": entry.filename, "c2": size, "c3": digest.hexdigest()})
    yield {"type": "warning", "reason": "ARCHIVE_MEMBERS_NOT_PARSED"}


def parse_asset(path, kind, maximum_records=MAX_RECORDS):
    path = Path(path)
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise ParseError("SIZE_LIMIT")
    if not isinstance(maximum_records, int) or not 1 <= maximum_records <= MAX_RECORDS:
        raise ParseError("RECORD_LIMIT")
    count = 0
    features = 0
    reason = None
    try:
        if kind == "xlsx":
            events = workbook(path)
        elif kind == "xls":
            events = legacy_workbook(path)
        elif kind in ("html", "md", "pdf", "txt"):
            events = document(path, kind)
        elif kind == "zip":
            events = archive_inventory(path)
        else:
            raise ParseError("FORMAT_UNSUPPORTED")
        for event in events:
            if event["type"] == "warning":
                reason = event["reason"]
                continue
            if event["type"] == "record":
                count += 1
                if count > maximum_records:
                    raise ParseError("RECORD_LIMIT")
                event["index"] = count
                features += int(event["geometry"] is not None)
            yield event
        yield {
            "type": "summary",
            "status": "PARTIAL" if reason else "READY" if count else "EMPTY",
            "recordCount": count,
            "featureCount": features,
            "reason": reason,
        }
    except ParseError:
        raise
    except (
        ValueError,
        UnicodeError,
        zipfile.BadZipFile,
        OSError,
        KeyError,
        EOFError,
    ) as error:
        raise ParseError("INVALID_CONTENT") from error
